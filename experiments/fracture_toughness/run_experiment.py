"""Reproduce and validate the fracture-toughness fuzzy surface.

All 26 cleaned experimental observations are used as the training set. Each
observation defines one zero-order Takagi-Sugeno rule whose consequent is the
observed positive KIC value; consequents are not refitted. The model predicts
surfaces at measured and intermediate porosity values. OpenFLL predictions are
compared with Simpful and an independent NumPy oracle. If a separately supplied
Niihara reference dataset is available, it is evaluated as an external test set.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import platform
import sys
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

import numpy as np


INPUT_NAMES = ["P", "a", "c"]
OUTPUT_NAME = "KIC"
DEFAULT_GRID_SIZE = 100
DEFAULT_SIMPFUL_SURFACE_SAMPLE = 300
ARTICLE_A_RANGE = (14.0, 39.0)
ARTICLE_C_RANGE = (48.0, 109.0)
P_SWEEP_STEP = 0.1
TARGET_CONCENTRATIONS = (0.0, 0.1, 0.2, 0.3, 0.4, 0.5)


@dataclass(frozen=True)
class Observation:
    record_id: int
    concentration_mass_pct: float
    P: float
    a: float
    c: float
    KIC: float


@dataclass(frozen=True)
class RuleDefinition:
    index: int
    p_term: str
    a_term: str
    c_term: str
    output_term: str
    output_value: float


@dataclass(frozen=True)
class ModelDefinition:
    p_terms: tuple[tuple[str, float, float], ...]
    a_terms: tuple[tuple[str, float, float], ...]
    c_terms: tuple[tuple[str, float, float], ...]
    rules: tuple[RuleDefinition, ...]


def load_observations(path: Path) -> list[Observation]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        records = [
            Observation(
                record_id=int(row["record_id"]),
                concentration_mass_pct=float(row["concentration_mass_pct"]),
                P=float(row["P"]),
                a=float(row["a"]),
                c=float(row["c"]),
                KIC=float(row["KIC"]),
            )
            for row in reader
        ]
    if len(records) != 26:
        raise ValueError(f"Expected the article-cleaned 26 observations, got {len(records)}")
    if len({item.record_id for item in records}) != len(records):
        raise ValueError("record_id values must be unique")
    return records


def _term_name(prefix: str, index: int) -> str:
    return f"{prefix}_{index:02d}"


def neighbor_sigmas(values: Sequence[float], *, minimum_span_fraction: float = 0.05) -> dict[float, float]:
    """Return deterministic Gaussian widths from sorted unique centers.

    For an isolated adjacent gap d, sigma=d/(2*sqrt(2*ln(2))) makes equal
    Gaussians reach membership 0.5 at their midpoint. For irregularly spaced
    centers the mean adjacent gap is used, with a lower bound proportional to
    the complete domain span to prevent numerical holes in the response surface.
    """
    centers = np.asarray(sorted(set(float(v) for v in values)), dtype=np.float64)
    if centers.size < 2:
        raise ValueError("At least two distinct centers are required")
    domain_span = float(centers[-1] - centers[0])
    min_sigma = max(domain_span * minimum_span_fraction, np.finfo(np.float64).eps)
    half_height_factor = 2.0 * math.sqrt(2.0 * math.log(2.0))
    result: dict[float, float] = {}
    for index, center in enumerate(centers):
        gaps: list[float] = []
        if index > 0:
            gaps.append(float(center - centers[index - 1]))
        if index + 1 < centers.size:
            gaps.append(float(centers[index + 1] - center))
        local_gap = float(np.mean(gaps))
        result[float(center)] = max(local_gap / half_height_factor, min_sigma)
    return result


def build_model_definition(records: Sequence[Observation]) -> ModelDefinition:
    p_values = [item.P for item in records]
    a_values = [item.a for item in records]
    c_values = [item.c for item in records]
    p_sigmas = neighbor_sigmas(p_values)
    a_sigmas = neighbor_sigmas(a_values)
    c_sigmas = neighbor_sigmas(c_values)

    sorted_p = sorted(set(p_values))
    p_name_by_value = {value: _term_name("p", index) for index, value in enumerate(sorted_p)}
    p_terms = tuple((p_name_by_value[value], value, p_sigmas[value]) for value in sorted_p)
    a_terms = tuple((_term_name("a", index), item.a, a_sigmas[item.a]) for index, item in enumerate(records))
    c_terms = tuple((_term_name("c", index), item.c, c_sigmas[item.c]) for index, item in enumerate(records))
    rules = tuple(
        RuleDefinition(
            index=index,
            p_term=p_name_by_value[item.P],
            a_term=_term_name("a", index),
            c_term=_term_name("c", index),
            output_term=_term_name("kic", index),
            output_value=item.KIC,
        )
        for index, item in enumerate(records)
    )
    return ModelDefinition(p_terms=p_terms, a_terms=a_terms, c_terms=c_terms, rules=rules)


def observations_to_arrays(records: Sequence[Observation]) -> tuple[np.ndarray, np.ndarray]:
    X = np.asarray([[item.P, item.a, item.c] for item in records], dtype=np.float64)
    y = np.asarray([item.KIC for item in records], dtype=np.float64)
    return X, y


def _gaussian(values: np.ndarray, centers: np.ndarray, sigmas: np.ndarray) -> np.ndarray:
    return np.exp(-0.5 * ((values[:, None] - centers[None, :]) / sigmas[None, :]) ** 2)


def predict_numpy(model: ModelDefinition, X: np.ndarray, *, consequents: np.ndarray | None = None) -> np.ndarray:
    X = np.asarray(X, dtype=np.float64)
    if X.ndim != 2 or X.shape[1] != 3:
        raise ValueError(f"X must have shape (n_samples, 3), got {X.shape}")

    p_term_by_name = {name: (center, sigma) for name, center, sigma in model.p_terms}
    p_centers = np.asarray([p_term_by_name[rule.p_term][0] for rule in model.rules])
    p_sigmas = np.asarray([p_term_by_name[rule.p_term][1] for rule in model.rules])
    a_centers = np.asarray([term[1] for term in model.a_terms])
    a_sigmas = np.asarray([term[2] for term in model.a_terms])
    c_centers = np.asarray([term[1] for term in model.c_terms])
    c_sigmas = np.asarray([term[2] for term in model.c_terms])
    if consequents is None:
        consequents = np.asarray([rule.output_value for rule in model.rules], dtype=np.float64)
    else:
        consequents = np.asarray(consequents, dtype=np.float64)

    weights = (
        _gaussian(X[:, 0], p_centers, p_sigmas)
        * _gaussian(X[:, 1], a_centers, a_sigmas)
        * _gaussian(X[:, 2], c_centers, c_sigmas)
    )
    totals = weights.sum(axis=1)
    result = np.full(X.shape[0], np.nan, dtype=np.float64)
    active = totals > 0.0
    result[active] = (weights[active] @ consequents) / totals[active]
    return result


def build_openfll_engine(model: ModelDefinition) -> Any:
    import openfll

    engine = openfll.SugenoEngine()
    engine.set_default_t_norm("prod_and")
    engine.add_input_var("P")
    engine.add_input_var("a")
    engine.add_input_var("c")
    engine.add_output_var(OUTPUT_NAME)
    for name, center, sigma in model.p_terms:
        engine.add_membership_func("P", name, "gaussian", [center, sigma])
    for name, center, sigma in model.a_terms:
        engine.add_membership_func("a", name, "gaussian", [center, sigma])
    for name, center, sigma in model.c_terms:
        engine.add_membership_func("c", name, "gaussian", [center, sigma])
    for rule in model.rules:
        engine.add_membership_func(OUTPUT_NAME, rule.output_term, "constant", [rule.output_value])
        engine.add_rule(
            antecedent=[("P", rule.p_term), ("a", rule.a_term), ("c", rule.c_term)],
            consequent=[(OUTPUT_NAME, rule.output_term)],
            t_norm="prod_and",
        )
    engine.build()
    return engine


def predict_openfll(engine: Any, X: np.ndarray) -> np.ndarray:
    return np.asarray(engine.predict_batch(np.asarray(X, dtype=np.float64), input_names=INPUT_NAMES), dtype=np.float64)


def build_simpful_system(model: ModelDefinition) -> Any:
    from simpful import FuzzySet, FuzzySystem, Gaussian_MF, LinguisticVariable

    system = FuzzySystem(operators=["AND_PRODUCT"], show_banner=False, verbose=False)

    def linguistic_variable(terms: Iterable[tuple[str, float, float]]) -> Any:
        terms = tuple(terms)
        fuzzy_sets = [
            FuzzySet(function=Gaussian_MF(mu=center, sigma=sigma), term=name)
            for name, center, sigma in terms
        ]
        centers = [term[1] for term in terms]
        sigmas = [term[2] for term in terms]
        universe = [min(c - 4.0 * s for c, s in zip(centers, sigmas)), max(c + 4.0 * s for c, s in zip(centers, sigmas))]
        return LinguisticVariable(fuzzy_sets, universe_of_discourse=universe)

    system.add_linguistic_variable("P", linguistic_variable(model.p_terms), verbose=False)
    system.add_linguistic_variable("a", linguistic_variable(model.a_terms), verbose=False)
    system.add_linguistic_variable("c", linguistic_variable(model.c_terms), verbose=False)
    for rule in model.rules:
        system.set_crisp_output_value(rule.output_term, rule.output_value, verbose=False)
    system.add_rules(
        [
            f"IF (P IS {rule.p_term}) AND (a IS {rule.a_term}) AND (c IS {rule.c_term}) "
            f"THEN ({OUTPUT_NAME} IS {rule.output_term})"
            for rule in model.rules
        ],
        verbose=False,
    )
    return system


def predict_simpful_from_strengths(system: Any, model: ModelDefinition, X: np.ndarray) -> np.ndarray:
    X = np.asarray(X, dtype=np.float64)
    strengths = np.asarray(
        system.get_firing_strengths(
            {"P": X[:, 0].tolist(), "a": X[:, 1].tolist(), "c": X[:, 2].tolist()}
        ),
        dtype=np.float64,
    )
    expected_shape = (len(X), len(model.rules))
    if strengths.shape != expected_shape and strengths.shape == (len(model.rules), len(X)):
        strengths = strengths.T
    if strengths.shape != expected_shape:
        raise RuntimeError(f"Unexpected Simpful firing-strength shape {strengths.shape}; expected {expected_shape}")
    totals = strengths.sum(axis=1)
    consequents = np.asarray([rule.output_value for rule in model.rules], dtype=np.float64)
    result = np.full(len(X), np.nan, dtype=np.float64)
    active = totals > 0.0
    result[active] = (strengths[active] @ consequents) / totals[active]
    return result


def predict_simpful_scalar(system: Any, X: np.ndarray) -> np.ndarray:
    values: list[float] = []
    for P, a, c in np.asarray(X, dtype=np.float64):
        system.set_variable("P", float(P), verbose=False)
        system.set_variable("a", float(a), verbose=False)
        system.set_variable("c", float(c), verbose=False)
        values.append(float(system.Sugeno_inference([OUTPUT_NAME], verbose=False)[OUTPUT_NAME]))
    return np.asarray(values, dtype=np.float64)


def metric_summary(y_true: np.ndarray, y_pred: np.ndarray) -> dict[str, float | int]:
    y_true = np.asarray(y_true, dtype=np.float64)
    y_pred = np.asarray(y_pred, dtype=np.float64)
    finite = np.isfinite(y_pred)
    if not finite.all():
        return {"n": int(len(y_true)), "finite": int(finite.sum()), "nan_or_inf": int((~finite).sum())}
    residual = y_pred - y_true
    mse = float(np.mean(residual**2))
    denominator = float(np.sum((y_true - np.mean(y_true)) ** 2))
    return {
        "n": int(len(y_true)),
        "finite": int(finite.sum()),
        "nan_or_inf": 0,
        "mse": mse,
        "rmse": math.sqrt(mse),
        "mae": float(np.mean(np.abs(residual))),
        "max_abs_error": float(np.max(np.abs(residual))),
        "r2": float(1.0 - np.sum(residual**2) / denominator),
    }


def metric_summary_by_porosity(
    X: np.ndarray,
    y_true: np.ndarray,
    y_pred: np.ndarray,
) -> list[dict[str, Any]]:
    """Return validation metrics independently for each theoretical P surface."""
    X = np.asarray(X, dtype=np.float64)
    y_true = np.asarray(y_true, dtype=np.float64)
    y_pred = np.asarray(y_pred, dtype=np.float64)
    if len(X) != len(y_true) or len(y_true) != len(y_pred):
        raise ValueError("Predictors, reference response and predictions must have equal lengths")
    rows: list[dict[str, Any]] = []
    for P in sorted(set(float(value) for value in X[:, 0])):
        mask = X[:, 0] == P
        rows.append({"P": P, **metric_summary(y_true[mask], y_pred[mask])})
    return rows


def parity_summary(left: np.ndarray, right: np.ndarray) -> dict[str, float | int]:
    left = np.asarray(left, dtype=np.float64)
    right = np.asarray(right, dtype=np.float64)
    finite = np.isfinite(left) & np.isfinite(right)
    difference = np.abs(left[finite] - right[finite])
    return {
        "n": int(len(left)),
        "finite_pairs": int(finite.sum()),
        "nonfinite_pairs": int((~finite).sum()),
        "max_abs_diff": float(np.max(difference)) if difference.size else math.nan,
        "mean_abs_diff": float(np.mean(difference)) if difference.size else math.nan,
        "rmse_diff": float(np.sqrt(np.mean(difference**2))) if difference.size else math.nan,
    }


def article_grid(grid_size: int) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    a_axis = np.linspace(*ARTICLE_A_RANGE, grid_size, dtype=np.float64)
    c_axis = np.linspace(*ARTICLE_C_RANGE, grid_size, dtype=np.float64)
    a_grid, c_grid = np.meshgrid(a_axis, c_axis, indexing="xy")
    return a_axis, c_axis, a_grid, c_grid


def concentration_to_porosity(
    records: Sequence[Observation], concentrations: Sequence[float] | np.ndarray
) -> np.ndarray:
    """Map concentration to porosity by piecewise-linear interpolation.

    The three measured pairs are preserved exactly. Intermediate concentrations
    are predictions inside the measured 0.0--0.5 mass-percent interval, not new
    experimental observations.
    """
    pairs = sorted({(item.concentration_mass_pct, item.P) for item in records})
    measured_concentrations = np.asarray([item[0] for item in pairs], dtype=np.float64)
    measured_porosity = np.asarray([item[1] for item in pairs], dtype=np.float64)
    requested = np.asarray(concentrations, dtype=np.float64)
    if np.any(requested < measured_concentrations[0]) or np.any(requested > measured_concentrations[-1]):
        raise ValueError("Concentration interpolation is restricted to the measured interval")
    return np.interp(requested, measured_concentrations, measured_porosity)


def create_surface_points(
    records: Sequence[Observation],
    grid_size: int,
    target_concentrations: Sequence[float] = TARGET_CONCENTRATIONS,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    a_axis, c_axis, a_grid, c_grid = article_grid(grid_size)
    requested_concentrations = np.asarray(target_concentrations, dtype=np.float64)
    porosities = concentration_to_porosity(records, requested_concentrations)
    blocks = []
    concentrations = []
    for concentration, P in zip(requested_concentrations, porosities):
        blocks.append(np.column_stack([np.full(a_grid.size, P), a_grid.ravel(), c_grid.ravel()]))
        concentrations.append(np.full(a_grid.size, concentration))
    return np.vstack(blocks), np.concatenate(concentrations), porosities, np.asarray([a_axis, c_axis])


def _validate_niihara_reference(
    X: np.ndarray,
    y: np.ndarray,
    *,
    source: str,
) -> tuple[np.ndarray, np.ndarray]:
    """Validate an external theoretical reference without changing row order."""
    X = np.asarray(X, dtype=np.float64)
    y = np.asarray(y, dtype=np.float64)
    if X.ndim != 2 or X.shape[1] != 3:
        raise ValueError(f"{source} predictors must have shape (n, 3), got {X.shape}")
    if y.ndim != 1:
        raise ValueError(f"{source} response must be one-dimensional, got {y.shape}")
    if len(X) == 0:
        raise ValueError(f"{source} dataset is empty")
    if len(X) != len(y):
        raise ValueError(
            f"{source} predictors and response have different row counts: {len(X)} and {len(y)}"
        )
    values = np.column_stack([X, y])
    if not np.all(np.isfinite(values)):
        raise ValueError(f"{source} dataset contains non-finite values")
    if np.any(values[:, 1:] <= 0.0):
        raise ValueError(f"{source} a, c and KIC values must be positive")
    if len(np.unique(X, axis=0)) != len(X):
        raise ValueError(f"{source} predictor rows must be unique")
    return X, y


def load_niihara_reference(path: Path) -> tuple[np.ndarray, np.ndarray]:
    """Load an external theoretical reference CSV with columns P, a, c and KIC."""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"P", "a", "c", "KIC"}
        missing = required.difference(reader.fieldnames or [])
        if missing:
            raise ValueError(f"Niihara reference is missing columns: {sorted(missing)}")
        rows = [
            [float(row["P"]), float(row["a"]), float(row["c"]), float(row["KIC"])]
            for row in reader
        ]
    values = np.asarray(rows, dtype=np.float64)
    if values.size == 0:
        X = np.empty((0, 3), dtype=np.float64)
        y = np.empty(0, dtype=np.float64)
    else:
        X = values[:, :3]
        y = values[:, 3]
    return _validate_niihara_reference(X, y, source="Niihara CSV reference")


def _xlsx_records(path: Path) -> tuple[list[str], list[tuple[Any, ...]]]:
    """Read the active worksheet as a header and data rows using openpyxl."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Reading the theoretical Excel workbooks requires openpyxl; "
            "install experiments/fracture_toughness/requirements.txt"
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(iterator)
        except StopIteration as error:
            raise ValueError(f"Excel workbook is empty: {path}") from error
        header = [str(value).strip() if value is not None else "" for value in raw_header]
        rows = [tuple(row) for row in iterator if any(value is not None for value in row)]
    finally:
        workbook.close()
    return header, rows


def load_niihara_workbooks(
    predictors_path: Path,
    response_path: Path,
) -> tuple[np.ndarray, np.ndarray]:
    """Load aligned theoretical predictors and Kc response from two Excel books."""
    predictor_header, predictor_rows = _xlsx_records(predictors_path)
    response_header, response_rows = _xlsx_records(response_path)

    predictor_columns = {name: index for index, name in enumerate(predictor_header)}
    missing_predictors = {"P", "a", "c"}.difference(predictor_columns)
    if missing_predictors:
        raise ValueError(
            f"Theoretical predictors workbook is missing columns: {sorted(missing_predictors)}"
        )
    response_columns = {name.casefold(): index for index, name in enumerate(response_header)}
    response_index = response_columns.get("kc", response_columns.get("kic"))
    if response_index is None:
        raise ValueError("Theoretical response workbook is missing the Kc or KIC column")
    if len(predictor_rows) != len(response_rows):
        raise ValueError(
            "Theoretical predictors and response workbooks have different row counts: "
            f"{len(predictor_rows)} and {len(response_rows)}"
        )

    X = np.asarray(
        [
            [
                float(row[predictor_columns["P"]]),
                float(row[predictor_columns["a"]]),
                float(row[predictor_columns["c"]]),
            ]
            for row in predictor_rows
        ],
        dtype=np.float64,
    )
    y = np.asarray([float(row[response_index]) for row in response_rows], dtype=np.float64)

    ratio_index = predictor_columns.get("c/a")
    if ratio_index is not None:
        supplied_ratio = np.asarray(
            [float(row[ratio_index]) for row in predictor_rows], dtype=np.float64
        )
        calculated_ratio = X[:, 2] / X[:, 1]
        if not np.all(np.isfinite(supplied_ratio)) or not np.allclose(
            supplied_ratio, calculated_ratio, rtol=1e-10, atol=1e-12
        ):
            raise ValueError("Theoretical predictors contain inconsistent c/a values")

    return _validate_niihara_reference(X, y, source="theoretical Excel reference")


def porosity_axis(records: Sequence[Observation], step: float = P_SWEEP_STEP) -> np.ndarray:
    if step <= 0.0:
        raise ValueError("Porosity step must be positive")
    minimum = min(item.P for item in records)
    maximum = max(item.P for item in records)
    values = minimum + np.arange(math.floor((maximum - minimum) / step) + 1, dtype=np.float64) * step
    if not math.isclose(float(values[-1]), maximum, rel_tol=0.0, abs_tol=1e-12):
        values = np.append(values, maximum)
    return values


def calculate_porosity_sweep(
    openfll_engine: Any,
    model: ModelDefinition,
    records: Sequence[Observation],
    grid_size: int,
    step: float = P_SWEEP_STEP,
) -> tuple[list[dict[str, float]], float]:
    """Calculate integral surface statistics for measured and unmeasured P."""
    _, _, a_grid, c_grid = article_grid(grid_size)
    rows: list[dict[str, float]] = []
    maximum_parity_difference = 0.0
    for P in porosity_axis(records, step):
        points = np.column_stack([np.full(a_grid.size, P), a_grid.ravel(), c_grid.ravel()])
        openfll_values = predict_openfll(openfll_engine, points)
        numpy_values = predict_numpy(model, points)
        maximum_parity_difference = max(
            maximum_parity_difference, float(np.max(np.abs(openfll_values - numpy_values)))
        )
        rows.append(
            {
                "P": float(P),
                "mean": float(np.mean(openfll_values)),
                "std": float(np.std(openfll_values)),
                "min": float(np.min(openfll_values)),
                "max": float(np.max(openfll_values)),
            }
        )
    return rows, maximum_parity_difference


def write_csv(path: Path, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def save_plots(
    results_dir: Path,
    surface_X: np.ndarray,
    concentrations: np.ndarray,
    openfll_values: np.ndarray,
    grid_size: int,
) -> list[str]:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return []

    written: list[str] = []
    unique_concentrations = sorted(set(float(v) for v in concentrations))
    for concentration in unique_concentrations:
        mask = concentrations == concentration
        a_grid = surface_X[mask, 1].reshape(grid_size, grid_size)
        c_grid = surface_X[mask, 2].reshape(grid_size, grid_size)
        z_openfll = openfll_values[mask].reshape(grid_size, grid_size)
        P = float(surface_X[mask, 0][0])

        fig = plt.figure(figsize=(8, 6))
        axis = fig.add_subplot(111, projection="3d")
        surface = axis.plot_surface(a_grid, c_grid, z_openfll, cmap="viridis", linewidth=0, antialiased=True)
        axis.set_xlabel("a (micrometers)")
        axis.set_ylabel("c (micrometers)")
        axis.set_zlabel("KIC (MPa*m^0.5)")
        measured = concentration in {0.0, 0.1, 0.5}
        status = "measured" if measured else "intermediate prediction"
        axis.set_title(
            f"OpenFLL fuzzy surface, MWCNT={concentration:.1f} mass.%\n"
            f"P={P:.5f} ({status})"
        )
        fig.colorbar(surface, ax=axis, shrink=0.65, pad=0.12)
        fig.tight_layout()
        surface_name = f"surface_openfll_{concentration:.1f}.png"
        fig.savefig(results_dir / surface_name, dpi=160)
        plt.close(fig)
        written.append(surface_name)
    return written


def save_porosity_sweep_plot(results_dir: Path, rows: Sequence[dict[str, float]]) -> str | None:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return None

    P = np.asarray([row["P"] for row in rows], dtype=np.float64)
    mean = np.asarray([row["mean"] for row in rows], dtype=np.float64)
    std = np.asarray([row["std"] for row in rows], dtype=np.float64)
    fig, axis = plt.subplots(figsize=(8, 5))
    axis.plot(P, mean, color="#1f77b4", linewidth=2, label="surface mean")
    axis.fill_between(P, mean - std, mean + std, color="#1f77b4", alpha=0.2, label="mean ± 1 std")
    axis.set_xlabel("Porosity P")
    axis.set_ylabel("KIC (MPa*m^0.5)")
    axis.set_title("Predicted integral fracture toughness across porosity")
    axis.grid(alpha=0.25)
    axis.legend()
    fig.tight_layout()
    name = "mean_kic_vs_porosity.png"
    fig.savefig(results_dir / name, dpi=160)
    plt.close(fig)
    return name


def save_niihara_validation_plots(
    results_dir: Path,
    X: np.ndarray,
    y_reference: np.ndarray,
    y_predicted: np.ndarray,
) -> list[str]:
    """Save parity and absolute-error maps for the external theoretical surfaces."""
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return []

    X = np.asarray(X, dtype=np.float64)
    y_reference = np.asarray(y_reference, dtype=np.float64)
    y_predicted = np.asarray(y_predicted, dtype=np.float64)
    residual = y_predicted - y_reference
    written: list[str] = []

    fig, axis = plt.subplots(figsize=(7, 6))
    for P in sorted(set(float(value) for value in X[:, 0])):
        mask = X[:, 0] == P
        axis.scatter(
            y_reference[mask],
            y_predicted[mask],
            s=8,
            alpha=0.28,
            label=f"P={P:g}",
        )
    lower = float(min(np.min(y_reference), np.min(y_predicted)))
    upper = float(max(np.max(y_reference), np.max(y_predicted)))
    axis.plot([lower, upper], [lower, upper], color="black", linewidth=1.2, linestyle="--")
    axis.set_xlabel("Theoretical KIC (MPa*m^0.5)")
    axis.set_ylabel("OpenFLL KIC (MPa*m^0.5)")
    axis.set_title("OpenFLL prediction against theoretical reference")
    axis.grid(alpha=0.2)
    axis.legend()
    fig.tight_layout()
    parity_name = "niihara_reference_parity.png"
    fig.savefig(results_dir / parity_name, dpi=170)
    plt.close(fig)
    written.append(parity_name)

    porosity_values = sorted(set(float(value) for value in X[:, 0]))
    fig, axes = plt.subplots(1, len(porosity_values), figsize=(6 * len(porosity_values), 5), squeeze=False)
    maximum_error = float(np.max(np.abs(residual)))
    color_plot = None
    for axis, P in zip(axes[0], porosity_values):
        mask = X[:, 0] == P
        a_values = np.unique(X[mask, 1])
        c_values = np.unique(X[mask, 2])
        error_grid = np.full((len(c_values), len(a_values)), np.nan, dtype=np.float64)
        a_indices = np.searchsorted(a_values, X[mask, 1])
        c_indices = np.searchsorted(c_values, X[mask, 2])
        error_grid[c_indices, a_indices] = np.abs(residual[mask])
        color_plot = axis.pcolormesh(
            a_values,
            c_values,
            error_grid,
            shading="auto",
            cmap="magma",
            vmin=0.0,
            vmax=maximum_error,
        )
        axis.set_xlabel("a (micrometers)")
        axis.set_ylabel("c (micrometers)")
        axis.set_title(f"Absolute error, P={P:g}")
    fig.subplots_adjust(left=0.055, right=0.90, bottom=0.12, top=0.88, wspace=0.28)
    if color_plot is not None:
        color_axis = fig.add_axes([0.92, 0.16, 0.012, 0.68])
        fig.colorbar(color_plot, cax=color_axis, label="Absolute KIC error")
    error_name = "niihara_absolute_error_maps.png"
    fig.savefig(results_dir / error_name, dpi=170)
    plt.close(fig)
    written.append(error_name)
    return written


def model_to_json(model: ModelDefinition) -> dict[str, Any]:
    return {
        "inputs": {
            "P": [{"name": n, "center": c, "sigma": s} for n, c, s in model.p_terms],
            "a": [{"name": n, "center": c, "sigma": s} for n, c, s in model.a_terms],
            "c": [{"name": n, "center": c, "sigma": s} for n, c, s in model.c_terms],
        },
        "rules": [asdict(rule) for rule in model.rules],
        "t_norm": "product",
        "defuzzification": "normalized_weighted_average",
    }


def stratified_surface_sample_indices(concentrations: np.ndarray, max_points: int) -> np.ndarray:
    if max_points < 1:
        raise ValueError("max_points must be positive")
    groups = sorted(set(float(value) for value in concentrations))
    base = max_points // len(groups)
    remainder = max_points % len(groups)
    selected: list[np.ndarray] = []
    for position, group in enumerate(groups):
        candidates = np.flatnonzero(concentrations == group)
        count = min(len(candidates), base + (1 if position < remainder else 0))
        offsets = np.linspace(0, len(candidates) - 1, count, dtype=int)
        selected.append(candidates[offsets])
    return np.sort(np.concatenate(selected))


def run_experiment(
    data_path: Path,
    results_dir: Path,
    grid_size: int = DEFAULT_GRID_SIZE,
    simpful_surface_sample: int = DEFAULT_SIMPFUL_SURFACE_SAMPLE,
    niihara_reference_path: Path | None = None,
    niihara_predictors_path: Path | None = None,
    niihara_response_path: Path | None = None,
) -> dict[str, Any]:
    if niihara_reference_path is not None and (
        niihara_predictors_path is not None or niihara_response_path is not None
    ):
        raise ValueError("Use either a combined Niihara CSV or the two Excel workbooks, not both")
    if (niihara_predictors_path is None) != (niihara_response_path is None):
        raise ValueError("Both theoretical predictors and response workbooks must be supplied")

    results_dir.mkdir(parents=True, exist_ok=True)
    records = load_observations(data_path)
    model = build_model_definition(records)
    X_train, y_train = observations_to_arrays(records)
    surface_X, concentrations, surface_porosity, _ = create_surface_points(records, grid_size)

    timings: dict[str, float] = {}
    t0 = time.perf_counter()
    openfll_engine = build_openfll_engine(model)
    openfll_train = predict_openfll(openfll_engine, X_train)
    openfll_surface = predict_openfll(openfll_engine, surface_X)
    timings["openfll_seconds"] = time.perf_counter() - t0

    t0 = time.perf_counter()
    numpy_train = predict_numpy(model, X_train)
    numpy_surface = predict_numpy(model, surface_X)
    timings["numpy_seconds"] = time.perf_counter() - t0

    t0 = time.perf_counter()
    simpful_system = build_simpful_system(model)
    simpful_train = predict_simpful_from_strengths(simpful_system, model, X_train)
    simpful_indices = stratified_surface_sample_indices(concentrations, simpful_surface_sample)
    simpful_surface = predict_simpful_from_strengths(simpful_system, model, surface_X[simpful_indices])
    simpful_scalar_check = predict_simpful_scalar(simpful_system, X_train[:5])
    timings["simpful_seconds"] = time.perf_counter() - t0

    niihara_validation: dict[str, Any]
    niihara_validation_plots: list[str] = []
    if niihara_reference_path is None and niihara_predictors_path is None:
        niihara_validation = {
            "status": "not_run_missing_external_reference",
            "accepted_inputs": {
                "combined_csv_columns": ["P", "a", "c", "KIC"],
                "excel_predictor_columns": ["P", "a", "c", "c/a (optional check)"],
                "excel_response_columns": ["Kc or KIC"],
            },
            "reason": (
                "No external theoretical reference was supplied. The 26 experimental records "
                "remain the training data and are not reused as an external test."
            ),
        }
    else:
        t0 = time.perf_counter()
        if niihara_reference_path is not None:
            niihara_X, niihara_y = load_niihara_reference(niihara_reference_path)
            reference_source: dict[str, Any] = {
                "format": "combined_csv",
                "path": str(niihara_reference_path),
            }
        else:
            assert niihara_predictors_path is not None
            assert niihara_response_path is not None
            niihara_X, niihara_y = load_niihara_workbooks(
                niihara_predictors_path,
                niihara_response_path,
            )
            reference_source = {
                "format": "paired_xlsx",
                "predictors_path": str(niihara_predictors_path),
                "response_path": str(niihara_response_path),
            }
        niihara_predictions = predict_openfll(openfll_engine, niihara_X)
        niihara_numpy_predictions = predict_numpy(model, niihara_X)
        metrics_by_porosity = metric_summary_by_porosity(
            niihara_X, niihara_y, niihara_predictions
        )
        timings["niihara_reference_seconds"] = time.perf_counter() - t0
        niihara_validation = {
            "status": "completed_external_reference",
            "role": "full theoretical reference; excluded from training and rule construction",
            "source": reference_source,
            "records": int(len(niihara_y)),
            "porosity_values": sorted(set(float(value) for value in niihara_X[:, 0])),
            "theoretical_KIC_range": {
                "min": float(np.min(niihara_y)),
                "max": float(np.max(niihara_y)),
            },
            "predicted_KIC_range": {
                "min": float(np.min(niihara_predictions)),
                "max": float(np.max(niihara_predictions)),
            },
            "metrics": metric_summary(niihara_y, niihara_predictions),
            "metrics_by_porosity": metrics_by_porosity,
            "implementation_parity_openfll_vs_numpy": parity_summary(
                niihara_predictions, niihara_numpy_predictions
            ),
        }
        write_csv(
            results_dir / "niihara_reference_predictions.csv",
            [
                "P", "a", "c", "KIC_theoretical_reference", "KIC_openfll",
                "residual_openfll_minus_reference", "abs_error",
            ],
            (
                [
                    *niihara_X[index], niihara_y[index], niihara_predictions[index],
                    niihara_predictions[index] - niihara_y[index],
                    abs(niihara_y[index] - niihara_predictions[index]),
                ]
                for index in range(len(niihara_y))
            ),
        )
        write_csv(
            results_dir / "niihara_metrics_by_porosity.csv",
            ["P", "n", "mse", "rmse", "mae", "max_abs_error", "r2"],
            (
                [
                    row["P"], row["n"], row["mse"], row["rmse"], row["mae"],
                    row["max_abs_error"], row["r2"],
                ]
                for row in metrics_by_porosity
            ),
        )
        niihara_validation_plots = save_niihara_validation_plots(
            results_dir, niihara_X, niihara_y, niihara_predictions
        )

    t0 = time.perf_counter()
    sweep_rows, sweep_openfll_numpy_max_abs_diff = calculate_porosity_sweep(
        openfll_engine, model, records, grid_size
    )
    timings["porosity_sweep_seconds"] = time.perf_counter() - t0

    write_csv(
        results_dir / "training_predictions.csv",
        [
            "record_id", "concentration_mass_pct", "P", "a", "c", "KIC_observed",
            "KIC_openfll_fixed", "KIC_simpful_fixed", "KIC_numpy_fixed",
            "abs_openfll_simpful", "abs_openfll_numpy",
        ],
        (
            [
                item.record_id, item.concentration_mass_pct, item.P, item.a, item.c, item.KIC,
                openfll_train[index], simpful_train[index], numpy_train[index],
                abs(openfll_train[index] - simpful_train[index]),
                abs(openfll_train[index] - numpy_train[index]),
            ]
            for index, item in enumerate(records)
        ),
    )
    write_csv(
        results_dir / "surface_predictions.csv",
        [
            "concentration_mass_pct", "P", "a", "c", "KIC_openfll_fixed",
            "KIC_numpy_fixed", "abs_openfll_numpy", "porosity_status",
        ],
        (
            [
                concentrations[index], *surface_X[index], openfll_surface[index],
                numpy_surface[index],
                abs(openfll_surface[index] - numpy_surface[index]),
                "measured" if concentrations[index] in {0.0, 0.1, 0.5} else "intermediate_prediction",
            ]
            for index in range(len(surface_X))
        ),
    )
    write_csv(
        results_dir / "simpful_surface_sample.csv",
        ["surface_index", "concentration_mass_pct", "P", "a", "c", "KIC_openfll", "KIC_simpful", "abs_diff"],
        (
            [
                int(index), concentrations[index], *surface_X[index], openfll_surface[index],
                simpful_surface[position], abs(openfll_surface[index] - simpful_surface[position]),
            ]
            for position, index in enumerate(simpful_indices)
        ),
    )
    write_csv(
        results_dir / "porosity_sweep.csv",
        ["P", "KIC_mean", "KIC_std", "KIC_min", "KIC_max"],
        ([row["P"], row["mean"], row["std"], row["min"], row["max"]] for row in sweep_rows),
    )

    plots = save_plots(results_dir, surface_X, concentrations, openfll_surface, grid_size)
    plots.extend(niihara_validation_plots)
    sweep_plot = save_porosity_sweep_plot(results_dir, sweep_rows)
    if sweep_plot is not None:
        plots.append(sweep_plot)
    surface_statistics = {
        f"{concentration:.1f}": {
            "P": float(surface_X[concentrations == concentration, 0][0]),
            "status": "measured" if concentration in {0.0, 0.1, 0.5} else "intermediate_prediction",
            "mean": float(np.mean(openfll_surface[concentrations == concentration])),
            "std": float(np.std(openfll_surface[concentrations == concentration])),
            "min": float(np.min(openfll_surface[concentrations == concentration])),
            "max": float(np.max(openfll_surface[concentrations == concentration])),
        }
        for concentration in sorted(set(float(value) for value in concentrations))
    }
    consequent_values = np.asarray([rule.output_value for rule in model.rules], dtype=np.float64)
    summary: dict[str, Any] = {
        "environment": {
            "python": sys.version,
            "platform": platform.platform(),
            "numpy": np.__version__,
            "openfll_engine_module": type(openfll_engine).__module__,
            "simpful": "2.12.0",
        },
        "dataset": {
            "records": len(records),
            "porosity_values": sorted({item.P for item in records}),
            "concentrations_mass_pct": sorted({item.concentration_mass_pct for item in records}),
        },
        "model": {
            "rules": len(model.rules),
            "p_terms": len(model.p_terms),
            "a_terms": len(model.a_terms),
            "c_terms": len(model.c_terms),
            "grid_size_per_surface": grid_size,
            "surface_points_total": len(surface_X),
            "surface_concentrations": [float(value) for value in TARGET_CONCENTRATIONS],
            "surface_porosity": [float(value) for value in surface_porosity],
            "porosity_sweep_step": P_SWEEP_STEP,
            "porosity_sweep_levels": len(sweep_rows),
            "simpful_surface_sample_points": int(len(simpful_indices)),
        },
        "consequents": {
            "source": "observed KIC values from all 26 training records; no consequent fitting",
            "count": int(len(consequent_values)),
            "min": float(np.min(consequent_values)),
            "max": float(np.max(consequent_values)),
            "all_positive": bool(np.all(consequent_values > 0.0)),
        },
        "fixed_model_training_metrics": metric_summary(y_train, openfll_train),
        "niihara_validation": niihara_validation,
        "parity": {
            "training_openfll_vs_simpful": parity_summary(openfll_train, simpful_train),
            "training_openfll_vs_numpy": parity_summary(openfll_train, numpy_train),
            "surface_sample_openfll_vs_simpful": parity_summary(openfll_surface[simpful_indices], simpful_surface),
            "surface_openfll_vs_numpy": parity_summary(openfll_surface, numpy_surface),
            "simpful_scalar_vs_strengths_first_5": parity_summary(simpful_scalar_check, simpful_train[:5]),
            "porosity_sweep_openfll_vs_numpy_max_abs_diff": sweep_openfll_numpy_max_abs_diff,
        },
        "surface_ranges": {
            "openfll_min": float(np.nanmin(openfll_surface)),
            "openfll_max": float(np.nanmax(openfll_surface)),
            "openfll_nonfinite": int((~np.isfinite(openfll_surface)).sum()),
            "simpful_sample_nonfinite": int((~np.isfinite(simpful_surface)).sum()),
        },
        "fixed_model_surface_statistics_by_concentration": surface_statistics,
        "porosity_sweep": {
            "P_min": float(sweep_rows[0]["P"]),
            "P_max": float(sweep_rows[-1]["P"]),
            "levels": len(sweep_rows),
            "mean_KIC_min": float(min(row["mean"] for row in sweep_rows)),
            "mean_KIC_max": float(max(row["mean"] for row in sweep_rows)),
        },
        "timings": timings,
        "plots": plots,
        "limitations": [
            "The supplied article does not include the optimized Gaussian centers and sigmas.",
            "This run uses deterministic neighbor-based initial widths, not the article's genetic optimization.",
            "All 26 experimental records are training data; training metrics are not reported as test accuracy.",
            "The theoretical reference is evaluated only after model construction and is never used to fit rules, membership functions or consequents.",
            "The external theoretical values validate agreement with the supplied analytical surface; they are not additional physical experiments.",
            "Intermediate-concentration surfaces are model predictions and have no independent experimental ground truth.",
            "Simpful surface parity uses a deterministic stratified sample because its rule evaluation is scalar and much slower than OpenFLL batch inference.",
        ],
    }
    (results_dir / "model_definition.json").write_text(
        json.dumps(model_to_json(model), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (results_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


def main() -> None:
    experiment_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, default=experiment_dir / "article_clean_26.csv")
    parser.add_argument("--results-dir", type=Path, default=experiment_dir / "results")
    parser.add_argument("--grid-size", type=int, default=DEFAULT_GRID_SIZE)
    parser.add_argument("--simpful-surface-sample", type=int, default=DEFAULT_SIMPFUL_SURFACE_SAMPLE)
    parser.add_argument(
        "--niihara-reference",
        type=Path,
        default=None,
        help="Optional external CSV test set with columns P,a,c,KIC",
    )
    parser.add_argument(
        "--niihara-predictors-xlsx",
        type=Path,
        default=None,
        help="Optional theoretical predictors workbook with columns a,c,P and optional c/a",
    )
    parser.add_argument(
        "--niihara-response-xlsx",
        type=Path,
        default=None,
        help="Optional theoretical response workbook with column Kc or KIC",
    )
    args = parser.parse_args()
    if args.grid_size < 2:
        parser.error("--grid-size must be at least 2")
    if args.simpful_surface_sample < 1:
        parser.error("--simpful-surface-sample must be positive")
    if args.niihara_reference is not None and (
        args.niihara_predictors_xlsx is not None or args.niihara_response_xlsx is not None
    ):
        parser.error("use either --niihara-reference or the two Excel options, not both")
    if (args.niihara_predictors_xlsx is None) != (args.niihara_response_xlsx is None):
        parser.error("both --niihara-predictors-xlsx and --niihara-response-xlsx are required")
    summary = run_experiment(
        args.data,
        args.results_dir,
        args.grid_size,
        args.simpful_surface_sample,
        args.niihara_reference,
        args.niihara_predictors_xlsx,
        args.niihara_response_xlsx,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
